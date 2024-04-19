import pandas as pd
import os
import re
import time

from openpyxl import load_workbook

dict1={}
dict2={}

# mapping_excel = '2024-03-25DG4278.xlsx'
mapping_excel = 'Book1.xlsx'
concat_excel = '2024-03-25_DG4278_Test_Matrix_NewReq.xlsx'

def main():
    start = time.time()
    wb1 = load_workbook(mapping_excel, read_only = False)
    wb2 = load_workbook(concat_excel, read_only = False)
    
    for sheet in range(1, len(wb1.sheetnames)):
        work1 = wb1[wb1.sheetnames[1]]
        work2 = wb2[wb2.sheetnames[1]]
        
        make_dict(work1, work2)
        ws = wb1.sheetnames[1]
        compare(work1, work2, ws)
        
    wb1.save('DG4278_Test_Matrix_NewReq.xlsx')
    end = time.time()
    print('Time elapsed: ' + str(start-end) + ' seconds')
    
def make_dict(work1, work2):
    for r in range(1, work1.max_column+1):
        key = work1.cell(1, r).value
        dict1[key] = []
        for c in range(2, work1.max_row+1):
            value = work1.cell(c, r).value
            dict1[key].append(value)
    # print(dict1)
    
    for r in range(1, work2.max_column+1):
        key = work2.cell(1, r).value
        dict2[key] = []
        for c in range(2, work2.max_row+1):
            value = work2.cell(c, r).value
            dict2[key].append(value)
    # print(dict2)                
    
def compare(work1, work2, ws):
    k=1 #判斷key在哪一列
    a=0
    for key in dict1:
        for i in range(1, work1.max_row):
            row_index = i+1+a
            if key == 'Key':
                for b, CPEGR in enumerate(dict2.get('Key')):            
                    if type(CPEGR) == str and is_blank_or_none(CPEGR) == False and  'CPEGR' in CPEGR and is_blank_or_none(dict1[key][i-1]) == False:
                        # print(dict1[key][i-1], '111', CPEGR, CPEGR.find('\n') != -1, re.search(dict1[key][i-1], CPEGR))
                        if re.search(dict1[key][i-1], CPEGR):
                            print(f'查看 {concat_excel} 中 {work1}，第 {b+2} 列有相似的 {dict1[key][i-1]}') #第一個excel 從1開始，第二個excel從0開始，整體少2，故+2  
                            # print(work1.cell(work1.cell(i+1 , k).row, 1).value) 
                            # print(work2.cell(b+2, 1).value) 
                            # print(work1.cell(i+1 , k).row)    #確認在哪一行
                            
                            # print(dict1[key][i-1], work1.cell(work1.cell(row_index, k).row, k).value)
                            print(i, check(b, work2, dict1[key][i-1]))
                            if check(b, work2, dict1[key][i-1]) > 0 and dict1[key][i-1] == work1.cell(work1.cell(row_index, k).row, k).value:
                                work1.insert_rows(work1.cell(row_index , k).row+1, check(b, work2, dict1[key][i-1]))
                                print(f'{work1.cell(work1.cell(row_index , k).row, k).value} 底下以新增 {check(b, work2, dict1[key][i-1])} 個空白行')
                                a+=check(b, work2, dict1[key][i-1])
                        
                            
                            work1.cell(work1.cell(row_index, k).row, 1).value = work2.cell(b+2, 1).value
                            work1.cell(work1.cell(row_index, k).row, 6).value = work2.cell(b+2, 6).value
                            
                            
            else:
                continue
                    
        k+=1
                    
def is_blank_or_none(s):
    return s is None or len(s.strip()) == 0    

#判斷某行底下幾個空白
def check(b, work2, CPEGR):
    c=0
    for row_index in range(b+2, work2.max_row):
        cell_value = work2.cell(row=row_index, column=2).value
        if not cell_value or str(cell_value).strip() == "":
            c+=1
        elif cell_value == CPEGR:
            continue
        else:
            break
    
    return c
    
    
if __name__ == "__main__":
    main()